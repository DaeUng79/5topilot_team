import re
from datetime import date, datetime
from io import BytesIO
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import msoffcrypto
from msoffcrypto.exceptions import DecryptionError, FileFormatError, InvalidKeyError
from openpyxl import load_workbook

from .errors import DomainError

HEADERS = [
    "연번",
    "소속",
    "직위",
    "성명",
    "생년월일",
    "근무 지역",
    "구분",
    "부서",
    "전문분야",
    "휴대전화",
    "비고",
    "1차 선정 여부",
    "최종 선정 여부",
    "발주사업명",
]
KEYS = [
    "source_serial",
    "organization",
    "position",
    "name",
    "birth_date",
    "region",
    "category",
    "department",
    "specialty",
    "phone",
    "memo",
    "legacy_first",
    "legacy_final",
    "legacy_project",
]
MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 5000


def fields_of(text: str) -> list[str]:
    parts = [part.strip() for part in re.split(r"[,，;；\n]+", text)]
    return sorted({"인공지능" if part.upper() == "AI" else part for part in parts if part})


def phone_of(text: str) -> str:
    if not text.strip():
        return ""
    phone = re.sub(r"[-\s()]", "", text)
    if phone.startswith("+82"):
        phone = "0" + phone[3:]
    if not re.fullmatch(r"01[016789]\d{7,8}", phone):
        raise DomainError(422, "휴대전화는 0으로 시작하는 휴대전화 번호를 입력하세요.")
    return phone


def text_of(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def email_of(text: str) -> str:
    value = text.strip().lower()
    if value and (len(value) > 254 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value)):
        raise DomainError(422, "이메일 주소를 확인하세요.")
    return value


def prepare_candidate(body: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {key: text_of(body.get(key, "")) for key in KEYS}
    if any(len(value) > 2000 for value in result.values()):
        raise DomainError(422, "입력 항목은 2,000자 이하로 제한됩니다.")
    result["phone"] = phone_of(result["phone"])
    result["email"] = email_of(text_of(body.get("email", "")))
    result["fields"] = fields_of(result["specialty"])
    result["enabled"] = bool(body.get("enabled", True))
    warnings = []
    for key, label in [("name", "성명"), ("organization", "소속"), ("specialty", "전문분야")]:
        if not result[key]:
            warnings.append(f"{label} 보완 필요")
    if not result["phone"]:
        warnings.append("휴대전화 보완 필요: 통화·인증 준비 안 됨")
    result["warnings"] = warnings
    result["active"] = bool(
        result["enabled"] and result["name"] and result["organization"] and result["fields"]
    )
    result["auth_ready"] = bool(result["active"] and (result["phone"] or result["email"]))
    return result


def read_candidates(
    raw: bytes, password: str, filename: str, *, strict: bool = True
) -> list[dict[str, Any]]:
    if not filename.lower().endswith(".xlsx") or not raw or len(raw) > MAX_BYTES:
        raise DomainError(422, "8MB 이하의 .xlsx 파일을 선택하세요.")
    stream = BytesIO(raw)
    if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        if not password:
            raise DomainError(422, "암호화된 엑셀입니다. 파일 열기 암호를 입력하세요.")
        try:
            office = msoffcrypto.OfficeFile(stream)
            office.load_key(password=password, verify_password=True)
            stream = BytesIO()
            office.decrypt(stream)
        except (DecryptionError, FileFormatError, InvalidKeyError, OSError, ValueError) as exc:
            raise DomainError(
                422, "파일 열기 암호가 맞지 않거나 암호 파일을 읽을 수 없습니다."
            ) from exc
        if stream.getbuffer().nbytes > 32 * 1024 * 1024:
            raise DomainError(422, "복호화된 파일이 허용 크기를 초과합니다.")
        stream.seek(0)
    try:
        with ZipFile(stream) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(item.file_size for item in items) > 64 * 1024 * 1024:
                raise DomainError(422, "엑셀 압축 해제 크기가 허용 한도를 초과합니다.")
            if any(item.filename.endswith("vbaProject.bin") for item in items):
                raise DomainError(422, "매크로가 포함된 파일은 등록할 수 없습니다.")
            for entry in items:
                if entry.filename.endswith(".xml"):
                    xml = archive.read(entry)
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise DomainError(422, "외부 엔터티가 포함된 XML은 등록할 수 없습니다.")
        stream.seek(0)
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, OSError, ValueError, KeyError, ParseError) as exc:
        raise DomainError(422, "손상되었거나 지원하지 않는 엑셀 파일입니다.") from exc
    try:
        result = []
        seen_phones: set[str] = set()
        header_keys = {
            re.sub(r"\s", "", label): key for label, key in zip(HEADERS, KEYS, strict=True)
        }
        header_keys["이메일"] = "email"
        for sheet in book.worksheets:
            if sheet.max_row and sheet.max_row > MAX_ROWS + 30:
                raise DomainError(422, "서식별 후보는 최대 5,000행까지 등록할 수 있습니다.")
            if sheet.max_column and sheet.max_column > 100:
                raise DomainError(422, "열 수가 허용 한도를 초과합니다.")
            mapping: dict[int, str] = {}
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                if not mapping:
                    labels = [re.sub(r"\s", "", text_of(cell.value)) for cell in row]
                    if {"성명", "소속", "전문분야", "휴대전화"}.issubset(labels):
                        mapping = {
                            i: header_keys[label]
                            for i, label in enumerate(labels)
                            if label in header_keys
                        }
                        if len(mapping) != len(set(mapping.values())):
                            raise DomainError(422, "중복된 서식 열 이름이 있습니다.")
                    elif row_number > 20:
                        break
                    continue
                if not any(cell.value is not None for cell in row):
                    continue
                if any(cell.data_type == "f" for cell in row):
                    raise DomainError(
                        422, f"{row_number}행에 수식이 있습니다. 값으로 바꿔 등록하세요."
                    )
                body = {key: text_of(row[index].value) for index, key in mapping.items()}
                if strict and not body.get("organization") and not body.get("name"):
                    continue
                if not strict:
                    result.append(body | {"source_row": row_number, "source_sheet": sheet.title})
                    if len(result) > MAX_ROWS:
                        raise DomainError(422, "후보는 최대 5,000행까지 등록할 수 있습니다.")
                    continue
                try:
                    item = prepare_candidate(body)
                except DomainError as exc:
                    raise DomainError(422, f"{row_number}행: {exc.message}") from exc
                if item["phone"] and item["phone"] in seen_phones:
                    raise DomainError(422, f"{row_number}행: 파일 내 휴대전화가 중복됩니다.")
                if item["phone"]:
                    seen_phones.add(item["phone"])
                item["source_row"] = row_number
                item["source_sheet"] = sheet.title
                result.append(item)
                if len(result) > MAX_ROWS:
                    raise DomainError(422, "후보는 최대 5,000행까지 등록할 수 있습니다.")
        if not result:
            raise DomainError(422, "부서 서식 열 또는 등록할 후보 행을 찾지 못했습니다.")
        return result
    finally:
        book.close()
