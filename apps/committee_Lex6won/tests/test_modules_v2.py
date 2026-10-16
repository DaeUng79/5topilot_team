"""Contracts for Lex6won's headless v2 modules; all identities are synthetic."""

import asyncio
import json
import secrets
import sqlite3
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import load_workbook

import candidate_import as imports
import candidate_pool as pool
import committee_roster as roster
import committee_selection as selection
from committee_app.v2_core import Context, ModuleError, Runtime, digest, now
from committee_app.v2_excel import export_xlsx
from committee_app.v2_gateway import dispatch_action, error_response
from committee_app.v2_handoff import acknowledge_event, pending_events
from committee_app.v2_legacy import migrate_legacy_pool


def run(fn: Any, rt: Runtime, command: dict, actor: str = "admin-a", key: str = "") -> dict:
    return asyncio.run(fn(Context(actor, key or str(uuid4())), command, runtime=rt))


@pytest.fixture
def runtime(tmp_path: Path) -> Runtime:
    rt = Runtime(
        tmp_path / "admin.sqlite3",
        pii_key=secrets.token_bytes(32),
        lookup_key=secrets.token_bytes(32),
        auth_key=secrets.token_bytes(32),
    )
    asyncio.run(rt.initialize(development_parents=True))
    with rt.connect(write=True) as c:
        for suffix in ["a", "b"]:
            rt.insert(
                c,
                "departments",
                {
                    "id": "dept-" + suffix,
                    "organization_name": "가상기관",
                    "name": suffix,
                    "code": suffix,
                    "active": 1,
                },
            )
            rt.insert(
                c,
                "admin_users",
                {
                    "id": "admin-" + suffix,
                    "department_id": "dept-" + suffix,
                    "username": suffix,
                    "password_hash": "fixture-not-a-password",
                    "role": "DEPARTMENT_ADMIN",
                    "display_name": "가상 관리자",
                    "active": 1,
                    "password_changed_at": now(),
                },
            )
            rt.insert(
                c,
                "projects",
                {
                    "id": "project-" + suffix,
                    "department_id": "dept-" + suffix,
                    "name": "가상 사업",
                    "contact_organization": "가상기관",
                    "contact_department": suffix,
                    "contact_position": "가상직책",
                    "contact_name": "가상 담당자",
                    "contact_phone_cipher": b"fixture",
                    "contact_email_cipher": None,
                    "status": "PREPARING",
                    "created_by": "admin-" + suffix,
                },
            )
        for sid in ["AI", "IT", "SEC"]:
            rt.insert(
                c,
                "specialties",
                {
                    "id": sid,
                    "code": sid,
                    "name": sid,
                    "aliases_json": json.dumps([sid, "인공지능"] if sid == "AI" else [sid]),
                    "active": 1,
                },
            )
    return rt


def candidate(i: int, specialties: list[str] | None = None, **extra: Any) -> dict:
    return {
        "department_id": "dept-a",
        "name": f"가상위원{i}",
        "organization_name": "가상회사",
        "phone": f"010{i:08d}",
        "specialty_ids": specialties or ["AI"],
        **extra,
    }


def rules(
    rt: Runtime,
    *,
    direct: bool = False,
    target: int = 7,
    reserves: int = 0,
    quotas: list[dict] | None = None,
) -> dict:
    return run(
        selection.save_rules,
        rt,
        {
            "project_id": "project-a",
            "selection_mode": "DIRECT" if direct else "POOL_RANDOM",
            "target_count": target,
            "reserve_count": reserves,
            "quotas": [] if direct else quotas or [{"specialty_id": "AI", "final_count": target}],
        },
    )


def accept(rt: Runtime, nominees: list[dict]) -> list[dict]:
    return [
        run(
            roster.save_contact,
            rt,
            {
                "project_id": "project-a",
                "nominee_id": n["id"],
                "version": n["version"],
                "contact_status": "AVAILABLE",
            },
        )
        for n in nominees
    ]


def test_schema_privacy_and_department_boundary(runtime: Runtime) -> None:
    row = run(pool.create_candidate, runtime, candidate(1))
    assert row["version"] == 1
    with runtime.connect() as c:
        stored = c.execute("SELECT * FROM evaluator_candidates").fetchone()
        assert b"01000000001" not in stored["phone_cipher"]
        assert c.execute("PRAGMA foreign_keys").fetchone()[0] == 1
        assert not c.execute("PRAGMA foreign_key_check").fetchall()
    with pytest.raises(ModuleError, match="부서"):
        run(pool.search_candidates, runtime, {"department_id": "dept-a"}, actor="admin-b")
    with pytest.raises(ModuleError):
        run(pool.create_candidate, runtime, candidate(1))


@pytest.mark.parametrize("target", [7, 8, 9])
def test_random_quota_and_no_duplicate(runtime: Runtime, target: int) -> None:
    for i in range(1, 3 * target + 3):
        run(pool.create_candidate, runtime, candidate(i, ["AI", "IT"]))
    rule = rules(
        runtime,
        target=target,
        quotas=[
            {"specialty_id": "AI", "final_count": 3},
            {"specialty_id": "IT", "final_count": target - 3},
        ],
    )
    result = run(
        selection.draw_shortlist,
        runtime,
        {"project_id": "project-a", "rule_version": rule["version"]},
    )
    assert len(result["members"]) == 3 * target
    assert len({m["candidate_id"] for m in result["members"]}) == 3 * target
    assert Counter(m["specialty_id"] for m in result["members"]) == {
        "AI": 9,
        "IT": 3 * (target - 3),
    }


def test_direct_roster_reserves_and_outbox(runtime: Runtime) -> None:
    rule = rules(runtime, direct=True, reserves=1)
    for i in range(1, 9):
        run(pool.create_candidate, runtime, candidate(i, scope="PROJECT", project_id="project-a"))
    assert run(pool.search_candidates, runtime, {"department_id": "dept-a"})["total"] == 0
    nominees = accept(
        runtime, run(roster.list_nominees, runtime, {"project_id": "project-a"})["items"]
    )
    result = run(
        roster.confirm_roster,
        runtime,
        {
            "project_id": "project-a",
            "rule_version": rule["version"],
            "members": [
                {
                    "nominee_id": n["id"],
                    "version": n["version"],
                    "member_kind": "PRIMARY" if i < 7 else "RESERVE",
                    "reserve_order": None if i < 7 else 1,
                }
                for i, n in enumerate(nominees)
            ],
        },
    )
    assert len(result["assignments"]) == 8
    assert result["sync_status"] == "PENDING"
    with runtime.connect() as c:
        payload = json.loads(c.execute("SELECT payload_json FROM sync_outbox").fetchone()[0])
        assert len(payload["assignments"]) == 8
        assert all("phone" not in a and "login_lookup_hmac" in a for a in payload["assignments"])
        assert not c.execute("PRAGMA foreign_key_check").fetchall()
    with pytest.raises(ModuleError):
        run(
            roster.save_contact,
            runtime,
            {
                "project_id": "project-a",
                "nominee_id": nominees[0]["id"],
                "version": nominees[0]["version"],
                "contact_status": "UNAVAILABLE",
            },
        )


def test_idempotency_and_optimistic_version(runtime: Runtime) -> None:
    row = run(pool.create_candidate, runtime, candidate(1), key="same")
    assert run(pool.create_candidate, runtime, candidate(1), key="same") == row
    with pytest.raises(ModuleError) as exc:
        run(pool.create_candidate, runtime, candidate(2), key="same")
    assert exc.value.code == "IDEMPOTENCY_CONFLICT"
    body = {k: v for k, v in candidate(1).items() if k != "department_id"}
    body.update(candidate_id=row["id"], version=1, name="수정가상위원")
    changed = run(pool.update_candidate, runtime, body)
    assert changed["version"] == 2
    with pytest.raises(ModuleError) as exc:
        run(pool.update_candidate, runtime, body)
    assert exc.value.code == "VERSION_CONFLICT"


def test_concurrent_draw_one_result(runtime: Runtime) -> None:
    for i in range(1, 23):
        run(pool.create_candidate, runtime, candidate(i))
    rule = rules(runtime)

    def draw(_: int) -> str:
        try:
            return run(
                selection.draw_shortlist,
                runtime,
                {"project_id": "project-a", "rule_version": rule["version"]},
                key="concurrent",
            )["id"]
        except ModuleError as exc:
            return exc.code

    with ThreadPoolExecutor(2) as workers:
        ids = list(workers.map(draw, [1, 2]))
    assert ids[0] == ids[1]
    with runtime.connect() as c:
        assert c.execute("SELECT count(*) FROM committee_draws").fetchone()[0] == 1


def test_legacy_schema_never_overwritten(tmp_path: Path) -> None:
    path = tmp_path / "legacy.db"
    with sqlite3.connect(path) as c:
        c.execute("CREATE TABLE candidates(id TEXT)")
    rt = Runtime(
        path,
        pii_key=secrets.token_bytes(32),
        lookup_key=secrets.token_bytes(32),
        auth_key=secrets.token_bytes(32),
    )
    with pytest.raises(ModuleError) as exc:
        asyncio.run(rt.initialize(development_parents=True))
    assert exc.value.code == "LEGACY_SCHEMA"


def preview(rt: Runtime, rows: list[list], key: str = "preview") -> dict:
    file = export_xlsx("candidates.xlsx", imports.HEADERS, rows)
    return asyncio.run(
        imports.preview_import(
            Context("admin-a", key),
            {"department_id": "dept-a", "filename": file.filename},
            raw=file.content,
            runtime=rt,
        )
    )


def test_import_preview_errors_no_partial_commit(runtime: Runtime) -> None:
    result = preview(
        runtime,
        [
            ["가상1", "기관", None, None, "AI", "01012345678", None],
            ["가상2", "기관", None, None, "UNKNOWN", "01012345679", None],
        ],
    )
    assert result["valid_count"] == 1 and result["errors"][0]["code"] == "SPECIALTY_UNKNOWN"
    assert result["items"][0]["phone"] == "***-****-5678"
    with pytest.raises(ModuleError) as exc:
        run(imports.commit_import, runtime, {"preview_id": result["id"], "version": 1})
    assert exc.value.code == "IMPORT_ERRORS"
    assert run(pool.search_candidates, runtime, {"department_id": "dept-a"})["total"] == 0


def test_import_commit_encryption_and_replay(runtime: Runtime) -> None:
    result = preview(
        runtime, [["가상1", "기관", None, None, "AI", "01012345678", "fake@example.invalid"]]
    )
    with runtime.connect() as c:
        row = c.execute("SELECT preview_cipher FROM candidate_imports").fetchone()
        assert b"01012345678" not in row[0] and b"example.invalid" not in row[0]
    command = {"preview_id": result["id"], "version": 1}
    done = run(imports.commit_import, runtime, command, key="commit")
    assert done["count"] == 1
    assert run(imports.commit_import, runtime, command, key="commit") == done
    with pytest.raises(ModuleError):
        run(imports.commit_import, runtime, command, actor="admin-b")
    assert (
        run(pool.search_candidates, runtime, {"department_id": "dept-a"})["items"][0]["email"]
        == "fake@example.invalid"
    )


def test_import_rechecks_duplicates_atomically(runtime: Runtime) -> None:
    result = preview(
        runtime,
        [
            ["가상1", "기관", None, None, "AI", "01000000001", None],
            ["가상2", "기관", None, None, "AI", "01000000002", None],
        ],
    )
    run(pool.create_candidate, runtime, candidate(2))
    with pytest.raises(ModuleError):
        run(imports.commit_import, runtime, {"preview_id": result["id"], "version": 1})
    assert run(pool.search_candidates, runtime, {"department_id": "dept-a"})["total"] == 1
    with runtime.connect() as c:
        assert c.execute("SELECT status FROM candidate_imports").fetchone()[0] == "PREVIEW"


def test_import_expiry_and_file_limit(runtime: Runtime) -> None:
    result = preview(runtime, [["가상1", "기관", None, None, "AI", "01000000001", None]])
    with runtime.connect(write=True) as c:
        c.execute("UPDATE candidate_imports SET expires_at='2000-01-01T00:00:00+00:00'")
    with pytest.raises(ModuleError) as exc:
        run(imports.commit_import, runtime, {"preview_id": result["id"], "version": 1})
    assert exc.value.code == "PREVIEW_EXPIRED"
    with pytest.raises(ModuleError) as exc:
        asyncio.run(
            imports.preview_import(
                Context("admin-a", "oversize"),
                {"department_id": "dept-a", "filename": "huge.xlsx"},
                raw=b"x" * (8 * 1024 * 1024 + 1),
                runtime=runtime,
            )
        )
    assert exc.value.code == "FILE_INVALID"


def test_exports_formula_safety_and_template(runtime: Runtime) -> None:
    result = run(imports.download_template, runtime, {})
    book = load_workbook(BytesIO(result.content))
    assert [c.value for c in book.active[1]] == imports.HEADERS
    book.close()
    file = export_xlsx("export.xlsx", ["값"], [["=HYPERLINK(1)"], [" @malicious"], ["ordinary"]])
    book = load_workbook(BytesIO(file.content))
    assert book.active["A2"].data_type == "s" and book.active["A2"].value.startswith("'")
    assert book.active["A3"].value.startswith("'")
    book.close()


def test_matching_handles_overlap_and_hall_shortage() -> None:
    quotas = [{"specialty_id": "AI", "final_count": 1}, {"specialty_id": "IT", "final_count": 1}]
    people = [{"id": str(i), "specialty_ids": ["AI", "IT"] if i < 3 else ["AI"]} for i in range(6)]
    for _ in range(20):
        result = selection.match_slots(people, quotas)
        assert len(set(cid for cid, _ in result)) == 6
        assert all(int(cid) < 3 for cid, sid in result if sid == "IT")
    with pytest.raises(ModuleError) as exc:
        selection.match_slots(people[:5], quotas)
    assert exc.value.code == "POOL_SHORTAGE"


def test_shortage_rolls_back_draw_and_request(runtime: Runtime) -> None:
    rule = rules(runtime)
    with pytest.raises(ModuleError):
        run(
            selection.draw_shortlist,
            runtime,
            {"project_id": "project-a", "rule_version": rule["version"]},
            key="short",
        )
    with runtime.connect() as c:
        assert c.execute("SELECT count(*) FROM committee_draws").fetchone()[0] == 0
        assert (
            c.execute(
                "SELECT count(*) FROM idempotency_records WHERE request_key='short'"
            ).fetchone()[0]
            == 0
        )


def direct_ready(rt: Runtime, *, collision: bool = False) -> tuple[dict, list[dict]]:
    rule = rules(rt, direct=True)
    for i in range(1, 8):
        extra = {"name": "같은이름", "phone": f"010{i:04d}1234"} if collision else {}
        run(
            pool.create_candidate,
            rt,
            candidate(i, scope="PROJECT", project_id="project-a", **extra),
        )
    return rule, accept(rt, run(roster.list_nominees, rt, {"project_id": "project-a"})["items"])


def confirm_command(rule: dict, nominees: list[dict]) -> dict:
    return {
        "project_id": "project-a",
        "rule_version": rule["version"],
        "members": [
            {"nominee_id": n["id"], "version": n["version"], "member_kind": "PRIMARY"}
            for n in nominees
        ],
    }


def test_roster_login_collision_atomic(runtime: Runtime) -> None:
    rule, nominees = direct_ready(runtime, collision=True)
    with pytest.raises(ModuleError) as exc:
        run(roster.confirm_roster, runtime, confirm_command(rule, nominees))
    assert exc.value.code == "LOGIN_COLLISION"
    with runtime.connect() as c:
        assert c.execute("SELECT count(*) FROM committee_rosters").fetchone()[0] == 0
        assert c.execute("SELECT count(*) FROM sync_outbox").fetchone()[0] == 0


@pytest.mark.parametrize("mutation", ["stale", "excluded", "inactive", "unavailable"])
def test_roster_validates_latest_nominees(runtime: Runtime, mutation: str) -> None:
    rule, nominees = direct_ready(runtime)
    first = nominees[0]
    if mutation == "excluded":
        changed = run(
            roster.exclude_nominee,
            runtime,
            {
                "project_id": "project-a",
                "nominee_id": first["id"],
                "version": first["version"],
                "reason": "이해충돌 확인",
            },
        )
        nominees[0] = changed
    elif mutation == "inactive":
        run(
            pool.deactivate_candidate,
            runtime,
            {"candidate_id": first["candidate_id"], "version": 1},
        )
    else:
        changed = run(
            roster.save_contact,
            runtime,
            {
                "project_id": "project-a",
                "nominee_id": first["id"],
                "version": first["version"],
                "contact_status": "UNAVAILABLE",
            },
        )
        if mutation == "unavailable":
            nominees[0] = changed
    with pytest.raises(ModuleError):
        run(roster.confirm_roster, runtime, confirm_command(rule, nominees))


def test_handoff_envelope_hash_ack_and_frozen_rows(runtime: Runtime) -> None:
    rule, nominees = direct_ready(runtime)
    result = run(roster.confirm_roster, runtime, confirm_command(rule, nominees))
    events = asyncio.run(pending_events(Context("admin-a"), runtime=runtime))
    assert len(events) == 1 and events[0]["contract_name"] == "roster.confirmed"
    event = events[0]
    assert event["payload_sha256"] == digest(event["payload"])
    assert event["payload"]["roster_sha256"] == digest(event["payload"]["assignments"])
    assert asyncio.run(pending_events(Context("admin-b"), runtime=runtime)) == []
    ack = {
        k: event[k] for k in ["event_id", "schema_version", "aggregate_version", "payload_sha256"]
    } | {"status": "APPLIED"}
    with pytest.raises(ModuleError):
        asyncio.run(
            acknowledge_event(Context("admin-a"), ack | {"payload_sha256": "bad"}, runtime=runtime)
        )
    done = asyncio.run(acknowledge_event(Context("admin-a"), ack, runtime=runtime))
    assert done["sync_status"] == "SENT"
    assert asyncio.run(pending_events(Context("admin-a"), runtime=runtime)) == []
    assert asyncio.run(acknowledge_event(Context("admin-a"), ack, runtime=runtime)) == done
    file = run(
        roster.export_roster, runtime, {"project_id": "project-a", "roster_id": result["id"]}
    )
    assert file.content[:2] == b"PK"
    with pytest.raises(sqlite3.IntegrityError), runtime.connect(write=True) as c:
        c.execute("UPDATE committee_assignments SET name_snapshot='tampered'")


def test_pool_search_update_reset_contacts_and_deactivate(runtime: Runtime) -> None:
    row = run(pool.create_candidate, runtime, candidate(1))
    rules(runtime, direct=True)
    n = run(roster.add_nominee, runtime, {"project_id": "project-a", "candidate_id": row["id"]})
    accept(runtime, [n])
    body = {k: v for k, v in candidate(1).items() if k != "department_id"}
    run(
        pool.update_candidate,
        runtime,
        body | {"candidate_id": row["id"], "version": 1, "name": "새 이름"},
    )
    nominees = run(roster.list_nominees, runtime, {"project_id": "project-a"})["items"]
    assert nominees[0]["contact_status"] == "NOT_CONTACTED"
    result = run(
        pool.search_candidates,
        runtime,
        {"department_id": "dept-a", "search": "새", "specialty_id": "AI", "active": True},
    )
    assert result["total"] == 1
    assert (
        run(pool.search_candidates, runtime, {"department_id": "dept-a", "search": "%"})["total"]
        == 0
    )
    run(pool.deactivate_candidate, runtime, {"candidate_id": row["id"], "version": 2})
    assert run(pool.search_candidates, runtime, {"department_id": "dept-a"})["total"] == 0
    assert (
        run(pool.search_candidates, runtime, {"department_id": "dept-a", "active": False})["total"]
        == 1
    )
    assert len(run(pool.list_specialties, runtime, {})["items"]) == 3


def test_rules_versions_constraints_and_cross_project(runtime: Runtime) -> None:
    rule = rules(runtime, direct=True)
    body = {
        "project_id": "project-a",
        "selection_mode": "DIRECT",
        "target_count": 8,
        "reserve_count": 0,
        "version": rule["version"],
    }
    assert run(selection.save_rules, runtime, body)["version"] == 2
    with pytest.raises(ModuleError):
        run(selection.save_rules, runtime, body)
    with pytest.raises(ModuleError):
        run(selection.save_rules, runtime, body | {"target_count": 6, "version": 2})
    row = run(pool.create_candidate, runtime, candidate(1))
    with pytest.raises(ModuleError):
        run(roster.add_nominee, runtime, {"project_id": "project-b", "candidate_id": row["id"]})


def test_cached_response_revalidates_department_and_account(runtime: Runtime) -> None:
    body = candidate(1, scope="PROJECT", project_id="project-a")
    rules(runtime, direct=True)
    run(pool.create_candidate, runtime, body, key="create")
    with runtime.connect(write=True) as c:
        c.execute("UPDATE projects SET department_id='dept-b' WHERE id='project-a'")
    with pytest.raises(ModuleError) as exc:
        run(pool.create_candidate, runtime, body, key="create")
    assert exc.value.code == "FORBIDDEN"
    with runtime.connect(write=True) as c:
        c.execute("UPDATE admin_users SET active=0 WHERE id='admin-a'")
    with pytest.raises(ModuleError) as exc:
        run(pool.list_specialties, runtime, {})
    assert exc.value.status == 401


def test_gateway_envelopes_and_client_identity_rejected(runtime: Runtime) -> None:
    ctx = Context("admin-a", "ui-request")
    result = asyncio.run(dispatch_action(ctx, "COM-05", candidate(1), runtime=runtime))
    assert result["meta"]["request_id"] == "ui-request"
    with pytest.raises(ModuleError) as exc:
        asyncio.run(
            dispatch_action(ctx, "COM-05", candidate(2) | {"actor_id": "admin-b"}, runtime=runtime)
        )
    assert error_response(ctx, exc.value)["error"]["code"] == "INPUT_INVALID"
    with pytest.raises(ModuleError):
        asyncio.run(dispatch_action(ctx, "SQL", {}, runtime=runtime))
    with pytest.raises(ModuleError):
        asyncio.run(dispatch_action(ctx, "COM-03", {}, runtime=runtime))
    file = asyncio.run(dispatch_action(ctx, "COM-02", {}, runtime=runtime))
    assert file.filename.endswith(".xlsx")


def test_legacy_explicit_mapping_preserves_id_and_source(runtime: Runtime, tmp_path: Path) -> None:
    source = tmp_path / "old.db"
    with sqlite3.connect(source) as c:
        c.execute(
            "CREATE TABLE candidates(id TEXT PRIMARY KEY,payload TEXT,"
            "active INTEGER,phone TEXT,version INTEGER)"
        )
        c.execute(
            "INSERT INTO candidates VALUES (?,?,?,?,?)",
            (
                "legacy-id",
                json.dumps({"name": "가상이전", "organization": "기관", "fields": ["인공지능"]}),
                1,
                "01000000001",
                0,
            ),
        )
    before = source.read_bytes()
    kwargs = {
        "source": source,
        "department_id": "dept-a",
        "specialty_map": {"인공지능": "AI"},
        "runtime": runtime,
    }
    report = asyncio.run(migrate_legacy_pool(Context("admin-a"), **kwargs))
    assert report["mode"] == "PREVIEW" and report["valid_count"] == 1
    done = asyncio.run(
        migrate_legacy_pool(
            Context("admin-a", "migration"), expected_sha256=report["source_sha256"], **kwargs
        )
    )
    assert done["imported_count"] == 1
    assert (
        run(pool.search_candidates, runtime, {"department_id": "dept-a"})["items"][0]["id"]
        == "legacy-id"
    )
    assert source.read_bytes() == before


def test_migration_reentrant_foreign_keys_and_wrong_key(runtime: Runtime) -> None:
    asyncio.run(runtime.initialize())
    row = run(pool.create_candidate, runtime, candidate(1))
    with runtime.connect() as c:
        assert c.execute("SELECT count(*) FROM schema_migrations").fetchone()[0] == 1
        cipher = c.execute(
            "SELECT phone_cipher FROM evaluator_candidates WHERE id=?", (row["id"],)
        ).fetchone()[0]
    wrong = Runtime(
        runtime.path,
        pii_key=secrets.token_bytes(32),
        lookup_key=secrets.token_bytes(32),
        auth_key=secrets.token_bytes(32),
    )
    with pytest.raises(ModuleError) as exc:
        wrong.decrypt(cipher, row["id"] + ":phone")
    assert exc.value.code == "KEY_OR_DATA_INVALID"
    with pytest.raises(sqlite3.IntegrityError), runtime.connect(write=True) as c:
        runtime.insert(
            c, "candidate_specialties", {"candidate_id": row["id"], "specialty_id": "missing"}
        )


def test_pool_roster_quota_and_download(runtime: Runtime) -> None:
    for i in range(1, 23):
        run(pool.create_candidate, runtime, candidate(i, ["AI", "IT"]))
    rule = rules(
        runtime,
        quotas=[{"specialty_id": "AI", "final_count": 3}, {"specialty_id": "IT", "final_count": 4}],
    )
    draw = run(selection.draw_shortlist, runtime, {"project_id": "project-a", "rule_version": 1})
    file = run(
        selection.export_shortlist, runtime, {"project_id": "project-a", "draw_id": draw["id"]}
    )
    book = load_workbook(BytesIO(file.content))
    assert book.active.max_row == 22
    book.close()
    nominees = accept(
        runtime, run(roster.list_nominees, runtime, {"project_id": "project-a"})["items"]
    )
    by_id = {n["id"]: n for n in nominees}
    grouped = {
        sid: [by_id[m["nominee_id"]] for m in draw["members"] if m["specialty_id"] == sid]
        for sid in ["AI", "IT"]
    }
    with pytest.raises(ModuleError) as exc:
        run(
            roster.confirm_roster,
            runtime,
            confirm_command(rule, grouped["AI"][:4] + grouped["IT"][:3]),
        )
    assert exc.value.code == "QUOTA_MISMATCH"
    result = run(
        roster.confirm_roster, runtime, confirm_command(rule, grouped["AI"][:3] + grouped["IT"][:4])
    )
    assert len(result["assignments"]) == 7
    assert all("login_lookup_hmac" not in item for item in result["assignments"])


def test_schema_fields_match_v2_contract(runtime: Runtime) -> None:
    contract = json.loads(
        (Path(__file__).parents[1] / "contracts/committee_schema_v2.json").read_text(
            encoding="utf8"
        )
    )
    with runtime.connect() as conn:
        for table in contract["tables"]:
            # Identifiers originate only from the checked-in schema, not request data.
            actual = {row[1] for row in conn.execute(f"PRAGMA table_info({table['name']})")}
            assert actual == {field[0] for field in table["fields"]}, table["name"]


def test_legacy_preview_reports_invalid_phone(runtime: Runtime, tmp_path: Path) -> None:
    source = tmp_path / "invalid-old.db"
    with sqlite3.connect(source) as conn:
        conn.execute(
            "CREATE TABLE candidates(id TEXT,payload TEXT,"
            "active INTEGER,phone TEXT,version INTEGER)"
        )
        conn.execute(
            "INSERT INTO candidates VALUES(?,?,?,?,?)",
            (
                "invalid",
                json.dumps({"name": "가상", "organization": "기관", "fields": ["AI"]}),
                1,
                "not-a-phone",
                0,
            ),
        )
    report = asyncio.run(
        migrate_legacy_pool(
            Context("admin-a"),
            source=source,
            department_id="dept-a",
            specialty_map={"AI": "AI"},
            runtime=runtime,
        )
    )
    assert report["valid_count"] == 0 and report["errors"][0]["code"] == "INPUT_INVALID"


def test_checked_in_handoff_fixture_and_input_contract() -> None:
    from committee_app.v2_contracts import CandidateCreate

    directory = Path(__file__).parents[1] / "contracts"
    event = json.loads((directory / "roster_confirmed_v2.fixture.json").read_text(encoding="utf8"))
    assert digest(event["payload"]) == event["payload_sha256"]
    assert digest(event["payload"]["assignments"]) == event["payload"]["roster_sha256"]
    assert len(event["payload"]["assignments"]) == 7
    assert event["contract_name"] == "roster.confirmed" and event["schema_version"] == 2
    for member in event["payload"]["assignments"]:
        assert member["project_id"] == event["project_id"]
        assert member["roster_id"] == event["aggregate_id"]
        assert not {"phone", "email", "nominee_id"}.intersection(member)
    schemas = json.loads((directory / "module_inputs_v2.json").read_text(encoding="utf8"))
    assert schemas["actions"]["COM-05"] == CandidateCreate.model_json_schema()


def test_incompatible_shared_schema_fails_before_migration(runtime: Runtime) -> None:
    with runtime.connect(write=True) as conn:
        conn.execute("ALTER TABLE projects RENAME COLUMN contact_name TO obsolete_contact")
    with pytest.raises(ModuleError) as exc:
        asyncio.run(runtime.initialize())
    assert exc.value.code == "SCHEMA_INCOMPATIBLE"
